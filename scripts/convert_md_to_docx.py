# -*- coding: utf-8 -*-
"""
scripts/convert_md_to_docx.py
Parses DomainAttackSurfaceScanner_Project_Report.md and converts it to a
professionally formatted Academic Word document:
DomainAttackSurfaceScanner_Project_Report.docx
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

MD_PATH = r"c:\Users\annap\Desktop\Projects\DomainAttackSurfaceScanner\DomainAttackSurfaceScanner_Project_Report.md"
DOCX_PATH = r"c:\Users\annap\Desktop\Projects\DomainAttackSurfaceScanner\DomainAttackSurfaceScanner_Project_Report.docx"

COLOR_NAVY = RGBColor(26, 54, 93)      # #1A365D
COLOR_SLATE = RGBColor(43, 108, 176)   # #2B6CB0
COLOR_STEEL = RGBColor(44, 82, 130)    # #2C5282
COLOR_BODY = RGBColor(45, 55, 72)      # #2D3748
COLOR_MUTED = RGBColor(113, 128, 150)  # #718096
COLOR_CODE = RGBColor(26, 32, 44)      # #1A202C

HEX_PRIMARY = "1A365D"
HEX_ROW_ALT = "F7FAFC"
HEX_BORDER = "CBD5E0"
HEX_CODE_BG = "F1F5F9"

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_run_text(p, text, is_bold=False, is_italic=False, color=COLOR_BODY, font_name="Calibri", font_size=11):
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.bold = is_bold
    run.italic = is_italic
    return run

def add_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=6, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    return p

def render_markdown_in_paragraph(p, text, font_size=11, base_color=COLOR_BODY):
    # Regex parser for bold (**text**) and code (`text`) and italics (*text*)
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            clean = token[2:-2]
            format_run_text(p, clean, is_bold=True, color=base_color, font_size=font_size)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            clean = token[1:-1]
            format_run_text(p, clean, font_name="Consolas", font_size=font_size - 1, color=COLOR_STEEL, is_bold=False)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            clean = token[1:-1]
            format_run_text(p, clean, is_italic=True, color=base_color, font_size=font_size)
        else:
            format_run_text(p, token, color=base_color, font_size=font_size)

def build_docx():
    print(f"Reading Markdown source: {MD_PATH}")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()

    # Configure Section: A4, 1-inch margins
    for s in doc.sections:
        s.page_width = Inches(8.27)   # A4 width
        s.page_height = Inches(11.69) # A4 height
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.different_first_page_header_footer = True

        # Header for normal pages
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Domain Attack Surface Scanner (DASS) — Technical Project Report")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        # Footer for normal pages
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Page ")
        frun.font.name = "Calibri"
        frun.font.size = Pt(9)
        frun.font.color.rgb = COLOR_MUTED
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        fp._p.append(fld)

    lines = md_text.splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        parsed_rows = []
        for line in tbl_lines:
            # strip start/end |
            trimmed = line.strip()
            if not trimmed:
                continue
            if re.match(r'^\|?\s*:?-+:?\s*(\|.*)*$', trimmed):
                # separator row, skip
                continue
            cells = [c.strip() for c in trimmed.strip('|').split('|')]
            parsed_rows.append(cells)
        if not parsed_rows:
            return

        col_count = max(len(r) for r in parsed_rows)
        # normalize all rows to col_count
        for r in parsed_rows:
            while len(r) < col_count:
                r.append("")

        tbl = doc.add_table(rows=len(parsed_rows), cols=col_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)

        for row_idx, row_data in enumerate(parsed_rows):
            row = tbl.rows[row_idx]
            is_header = (row_idx == 0)
            for col_idx, cell_value in enumerate(row_data):
                cell = row.cells[col_idx]
                set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15

                # Clean cell value (<br> replaced by newline)
                clean_val = cell_value.replace("<br>", "\n").replace("•", "▪")

                if is_header:
                    set_cell_background(cell, HEX_PRIMARY)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    render_markdown_in_paragraph(p, clean_val, font_size=9.5, base_color=RGBColor(255, 255, 255))
                    for run in p.runs:
                        run.bold = True
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "FFFFFF")
                    else:
                        set_cell_background(cell, HEX_ROW_ALT)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    render_markdown_in_paragraph(p, clean_val, font_size=9.5, base_color=COLOR_BODY)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def flush_code(lines):
        if not lines:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        set_cell_background(cell, HEX_CODE_BG)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        joined = "\n".join(lines)
        format_run_text(p, joined, font_name="Consolas", font_size=8.5, color=COLOR_CODE)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    i = 0
    total_lines = len(lines)
    while i < total_lines:
        line = lines[i]

        # Code block fence
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table rows
        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        stripped = line.strip()

        # Horizontal rule
        if stripped in ["---", "***", "___"]:
            # Subtle section divider
            p = add_styled_paragraph(doc, space_before=6, space_after=6)
            run = p.add_run("―" * 55)
            run.font.color.rgb = RGBColor(226, 232, 240)
            run.font.size = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, space_before=16, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
            format_run_text(p, h_text, is_bold=True, color=COLOR_NAVY, font_size=20, font_name="Calibri")
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        elif stripped.startswith("## "):
            h_text = stripped[3:].strip()
            p = add_styled_paragraph(doc, space_before=14, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
            format_run_text(p, h_text, is_bold=True, color=COLOR_SLATE, font_size=15, font_name="Calibri")
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        elif stripped.startswith("### "):
            h_text = stripped[4:].strip()
            p = add_styled_paragraph(doc, space_before=10, space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
            format_run_text(p, h_text, is_bold=True, color=COLOR_STEEL, font_size=12.5, font_name="Calibri")
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        elif stripped.startswith("#### "):
            h_text = stripped[5:].strip()
            p = add_styled_paragraph(doc, space_before=8, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
            format_run_text(p, h_text, is_bold=True, color=COLOR_BODY, font_size=11, font_name="Calibri")
            p.paragraph_format.keep_with_next = True
            i += 1
            continue

        # Bullet list items
        if stripped.startswith("- ") or stripped.startswith("* "):
            item_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, space_before=1, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.left_indent = Inches(0.25)
            format_run_text(p, "▪  ", is_bold=True, color=COLOR_SLATE, font_size=9)
            render_markdown_in_paragraph(p, item_text, font_size=10.5)
            i += 1
            continue

        # Numbered list items
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num = num_match.group(1)
            item_text = num_match.group(2)
            p = add_styled_paragraph(doc, space_before=1, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.left_indent = Inches(0.25)
            format_run_text(p, f"{num}.  ", is_bold=True, color=COLOR_SLATE, font_size=10.5)
            render_markdown_in_paragraph(p, item_text, font_size=10.5)
            i += 1
            continue

        # Normal paragraph or empty line
        if not stripped:
            i += 1
            continue

        # Standard paragraph
        p = add_styled_paragraph(doc, space_before=2, space_after=6, line_spacing=1.25, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        render_markdown_in_paragraph(p, stripped, font_size=11)
        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    print(f"Saving compiled DOCX to: {DOCX_PATH}")
    doc.save(DOCX_PATH)
    print(f"Successfully created: {DOCX_PATH}")

if __name__ == "__main__":
    build_docx()
